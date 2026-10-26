"""The Excel workbook people actually read. apps.json and contacts.json stay the machine-readable outputs."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import CATEGORIES, COUNTRIES
from .verify import best_per_type

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
STATUS_FILL = {"verified": PatternFill("solid", fgColor="D9EAD3"),
               "likely": PatternFill("solid", fgColor="FFF2CC"),
               "rejected": PatternFill("solid", fgColor="F4CCCC")}
LINK_FONT = Font(name=FONT, size=10, color="1155CC", underline="single")
BODY_FONT = Font(name=FONT, size=10)
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
CONTACT_TYPES = [("email", "Email"), ("x", "X"), ("linkedin", "LinkedIn"), ("threads", "Threads")]

# (header, width, value function). Value functions return a value, or (value, hyperlink).
APP_COLUMNS = [
    ("#", 5, lambda a, i, c: i),
    ("App", 34, lambda a, i, c: (a["name"], a.get("app_url"))),
    ("Developer (App Store seller)", 30, lambda a, i, c: a.get("seller_name") or a["publisher_name"]),
    ("Account", 11, lambda a, i, c: "Individual" if a.get("individual_account") else "Company"),
    ("Trend score", 9, lambda a, i, c: a["trend_score"]),
    ("Markets", 22, lambda a, i, c: " ".join(a["markets"])),
    ("Best rank", 8, lambda a, i, c: a["best_rank"]),
    ("Days charting", 9, lambda a, i, c: a.get("days_charting")),
    ("Rank change", 9, lambda a, i, c: a.get("rank_change")),
    ("Categories", 24, lambda a, i, c: ", ".join(x for x in a["categories"] if x != "Overall") or "Overall"),
    ("Released", 11, lambda a, i, c: _date(a.get("release_date"))),
    ("Apps by developer", 9, lambda a, i, c: a.get("portfolio_size")),
    ("Revenue / month (WW)", 14, lambda a, i, c: a.get("monthly_revenue")),
    ("Downloads / month (WW)", 14, lambda a, i, c: a.get("monthly_downloads")),
]
CONTACT_WIDTH = 28
VERIFIED_COL = get_column_letter(len(APP_COLUMNS) + len(CONTACT_TYPES) + 1)


def write_workbook(ranked: list[dict], contacts: dict, path: Path, dates: list[str]) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    founders = [a for a in ranked if a["tier"] == "founder"]
    small = [a for a in ranked if a["tier"] == "small_company"]
    excluded = [a for a in ranked if a["tier"] == "excluded"]
    _app_sheet(wb.create_sheet("Founders"), founders, contacts)
    _app_sheet(wb.create_sheet("Small companies"), small, contacts)
    _contacts_sheet(wb.create_sheet("Contacts"), ranked, contacts)
    _excluded_sheet(wb.create_sheet("Excluded"), excluded)
    _summary_sheet(summary, dates, bool(contacts))
    wb.save(path)


def _app_sheet(ws, apps: list[dict], contacts: dict) -> None:
    headers = [h for h, _, _ in APP_COLUMNS] + [label for _, label in CONTACT_TYPES] + \
              ["Verified contacts", "Website", "Why it passed"]
    widths = [w for _, w, _ in APP_COLUMNS] + [CONTACT_WIDTH] * len(CONTACT_TYPES) + [10, 30, 40]
    _header(ws, headers, widths)

    for i, app in enumerate(apps, start=1):
        found = contacts.get(str(app["app_id"]), {}).get("contacts", [])
        best = best_per_type(found)
        row = [fn(app, i, contacts) for _, _, fn in APP_COLUMNS]
        for kind, _ in CONTACT_TYPES:
            row.append((best[kind]["value"], best[kind]["url"], best[kind]["status"]) if kind in best else None)
        website = next((u["url"] for u in contacts.get(str(app["app_id"]), {}).get("official_urls", [])
                        if u["kind"] == "website"), app.get("seller_url"))
        searched = str(app["app_id"]) in contacts
        row += [sum(c["status"] == "verified" for c in found) if searched else None,
                (website, website), "; ".join(app["reasons"])]
        _write_row(ws, i + 1, row)

    _finish(ws, len(apps), score_col=headers.index("Trend score") + 1)
    for r in range(2, len(apps) + 2):
        ws.cell(r, headers.index("Released") + 1).number_format = "yyyy-mm-dd"
        ws.cell(r, headers.index("Trend score") + 1).number_format = "0.00"
        ws.cell(r, headers.index("Rank change") + 1).number_format = "+0;-0;0"


def _contacts_sheet(ws, ranked: list[dict], contacts: dict) -> None:
    headers = ["App", "Developer", "Tier", "Type", "Contact", "Status", "Why", "Checks", "Found on", "Pages"]
    _header(ws, headers, [30, 28, 13, 9, 32, 10, 46, 40, 44, 7])
    by_id = {str(a["app_id"]): a for a in ranked}
    order = {"verified": 0, "likely": 1, "rejected": 2}
    r = 1
    for app_id, result in contacts.items():
        app = by_id.get(app_id)
        if not app:
            continue
        for c in sorted(result["contacts"], key=lambda c: (order[c["status"]], c["type"])):
            r += 1
            checks = ", ".join(f"{k}: {_fmt_check(v)}" for k, v in c["checks"].items())
            _write_row(ws, r, [
                (app["name"], app.get("app_url")), app.get("seller_name") or app["publisher_name"],
                app["tier"].replace("_", " "), dict(CONTACT_TYPES)[c["type"]],
                (c["value"], c["url"], c["status"]), c["status"], c["note"], checks,
                (c["found_on"][0], c["found_on"][0]), len(c["found_on"]),
            ])
            ws.cell(r, 6).fill = STATUS_FILL[c["status"]]
    _finish(ws, r - 1)


def _excluded_sheet(ws, apps: list[dict]) -> None:
    headers = ["#", "App", "Publisher", "Trend score", "Markets", "Revenue / month (WW)", "Why it was excluded"]
    _header(ws, headers, [5, 34, 30, 9, 22, 14, 80])
    for i, a in enumerate(apps, start=1):
        _write_row(ws, i + 1, [i, (a["name"], a.get("app_url")), a["publisher_name"], a["trend_score"],
                               " ".join(a["markets"]), a.get("monthly_revenue"), "; ".join(a["reasons"])])
        ws.cell(i + 1, 4).number_format = "0.00"
    _finish(ws, len(apps))


def _summary_sheet(ws, dates: list[str], has_contacts: bool) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 80
    ws["A1"] = "App Scout"
    ws["A1"].font = Font(name=FONT, size=16, bold=True)
    span = dates[0] if len(dates) == 1 else f"{dates[-1]} to {dates[0]} ({len(dates)} days)"
    ws["A2"] = f"iOS App Store charts, {span}"
    ws["A2"].font = Font(name=FONT, size=10, color="666666")

    rows = [
        ("Markets", ", ".join(COUNTRIES), None),
        ("Categories", f"{len(CATEGORIES)}", ", ".join(CATEGORIES.values())),
        (None, None, None),
        ("Founder-built apps", "=COUNTA(Founders!B:B)-1", "Individual developer accounts, or company accounts with ≤5 apps and ≤3 years on the store"),
        ("Small companies", "=COUNTA('Small companies'!B:B)-1", "Passed the exclusions but older/larger company accounts, or new accounts already at $500k+/month"),
        ("Excluded", "=COUNTA(Excluded!B:B)-1", "Large publishers, $1M+/month revenue, apps over 4 years old, studio portfolios, AI chat wrappers"),
    ]
    if has_contacts:
        rows += [
            (None, None, None),
            ("Contacts verified", '=COUNTIF(Contacts!F:F,"verified")', "Every check passed and the contact ties back to the developer's own domain or name"),
            ("Contacts likely", '=COUNTIF(Contacts!F:F,"likely")', "On an official page and not contradicted, but a check couldn't confirm it (e.g. Threads needs login)"),
            ("Contacts rejected", '=COUNTIF(Contacts!F:F,"rejected")', "Failed a check: no mail server, profile doesn't exist, or unrelated account. Not shown on Founders"),
            ("Founder apps with a verified contact", f'=COUNTIF(Founders!{VERIFIED_COL}:{VERIFIED_COL},">0")', "Founder apps with at least one verified contact"),
        ]
    for i, (label, value, note) in enumerate(rows, start=4):
        for col, v in enumerate((label, value, note), start=1):
            cell = ws.cell(i, col, v)
            cell.font = Font(name=FONT, size=10, bold=col == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=col == 3)
        if isinstance(value, str) and value.startswith("="):
            ws.cell(i, 2).number_format = "#,##0"
        ws.cell(i, 2).alignment = Alignment(horizontal="right", vertical="top")

    legend = len(rows) + 6
    ws.cell(legend, 1, "Contact colours").font = Font(name=FONT, size=10, bold=True)
    for j, (status, text) in enumerate([("verified", "Verified"), ("likely", "Likely: check before use"),
                                        ("rejected", "Rejected (Contacts sheet only)")], start=1):
        cell = ws.cell(legend + j, 1, text)
        cell.font = BODY_FONT
        cell.fill = STATUS_FILL[status]
    sources = legend + 5
    ws.cell(sources, 1, "Sources").font = Font(name=FONT, size=10, bold=True)
    for j, text in enumerate([
        "Charts: Sensor Tower public category rankings (top 25 free/grossing/paid) and Apple top-chart RSS (top 100).",
        "Developer details: iTunes Lookup API (seller name, release dates, full app portfolio).",
        "Revenue and downloads are Sensor Tower's worldwide monthly estimate ranges, not exact figures.",
        "Contacts: only from the website, support and privacy-policy pages the developer lists on its App Store or Sensor Tower entry.",
    ], start=1):
        ws.cell(sources + j, 1, text).font = BODY_FONT
    ws.sheet_view.showGridLines = False


def _header(ws, headers: list[str], widths: list[int]) -> None:
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, col, h)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30


def _write_row(ws, r: int, values: list) -> None:
    for col, v in enumerate(values, start=1):
        status = None
        link = None
        if isinstance(v, tuple):
            v, link, *rest = v
            status = rest[0] if rest else None
        cell = ws.cell(r, col, v)
        cell.font = LINK_FONT if link and v else BODY_FONT
        if link and v:
            cell.hyperlink = link
        if status:
            cell.fill = STATUS_FILL[status]
        cell.border = THIN
        cell.alignment = Alignment(vertical="top")


def _finish(ws, n_rows: int, score_col: int | None = None) -> None:
    ws.freeze_panes = "C2"
    last = get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A1:{last}{max(n_rows + 1, 2)}"
    if score_col and n_rows:
        col = get_column_letter(score_col)
        ws.conditional_formatting.add(f"{col}2:{col}{n_rows + 1}", ColorScaleRule(
            start_type="min", start_color="FFFFFF", end_type="max", end_color="6AA84F"))


def _date(iso: str | None) -> date | None:
    return date.fromisoformat(iso[:10]) if iso else None


def _fmt_check(v) -> str:
    return {True: "pass", False: "fail", None: "can't check"}.get(v, str(v)) if not isinstance(v, str) else v
